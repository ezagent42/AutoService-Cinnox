"""
Data import utilities.
Supports importing from DOCX, PDF, and XLSX files.
"""

import json
import sys
from datetime import datetime
from pathlib import Path
from typing import Optional

from socialware.core import generate_id, ensure_dir
from socialware.config import get_domain_config


def extract_from_docx(file_path: Path) -> list[dict]:
    """Extract structured data from DOCX file."""
    try:
        from docx import Document
    except ImportError:
        print("Installing python-docx...")
        import subprocess
        subprocess.run([sys.executable, "-m", "pip", "install", "python-docx"], check=True)
        from docx import Document

    doc = Document(file_path)
    items = []
    current_item = {}

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            if current_item:
                items.append(current_item)
                current_item = {}
            continue

        if ':' in text:
            key, value = text.split(':', 1)
            current_item[key.strip().lower()] = value.strip()
        elif not current_item:
            current_item['name'] = text

    if current_item:
        items.append(current_item)

    return items


def extract_from_xlsx(file_path: Path) -> list[dict]:
    """Extract structured data from XLSX file."""
    try:
        import openpyxl
    except ImportError:
        print("Installing openpyxl...")
        import subprocess
        subprocess.run([sys.executable, "-m", "pip", "install", "openpyxl"], check=True)
        import openpyxl

    wb = openpyxl.load_workbook(file_path)
    ws = wb.active

    items = []
    headers = []

    for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
        if row_idx == 0:
            headers = [str(cell).lower().strip() if cell else f"col_{i}" for i, cell in enumerate(row)]
            continue

        if all(cell is None for cell in row):
            continue

        item = {}
        for i, cell in enumerate(row):
            if i < len(headers) and cell is not None:
                item[headers[i]] = str(cell)

        if item:
            items.append(item)

    return items


def extract_from_pdf(file_path: Path) -> list[dict]:
    """Extract structured data from PDF file."""
    try:
        import pdfplumber
    except ImportError:
        print("Installing pdfplumber...")
        import subprocess
        subprocess.run([sys.executable, "-m", "pip", "install", "pdfplumber"], check=True)
        import pdfplumber

    items = []
    current_item = {}

    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            for line in text.split('\n'):
                line = line.strip()
                if not line:
                    if current_item:
                        items.append(current_item)
                        current_item = {}
                    continue

                if ':' in line:
                    key, value = line.split(':', 1)
                    current_item[key.strip().lower()] = value.strip()
                elif not current_item:
                    current_item['name'] = line

    if current_item:
        items.append(current_item)

    return items


def save_item(item: dict, output_dir: Path, item_type: str) -> Path:
    """Save an item to its own folder."""
    name = item.get('name', item.get('product', item.get('customer', f'unnamed_{generate_id("")}')))
    item_id = generate_id(name)
    folder_name = f"{item_id}_{name[:30].replace(' ', '_').replace('/', '_')}"

    item_dir = ensure_dir(output_dir / folder_name)

    item['_id'] = item_id
    item['_type'] = item_type
    item['_created'] = datetime.now().isoformat()

    with open(item_dir / 'info.json', 'w', encoding='utf-8') as f:
        json.dump(item, f, ensure_ascii=False, indent=2)

    with open(item_dir / 'README.md', 'w', encoding='utf-8') as f:
        f.write(f"# {name}\n\n")
        for key, value in item.items():
            if not key.startswith('_'):
                f.write(f"**{key.title()}**: {value}\n\n")

    return item_dir


def import_file(file_path: str, output_dir: str, item_type: str) -> list[Path]:
    """Import data from a file and save to output directory."""
    file_path = Path(file_path)
    output_dir = Path(output_dir)

    if not file_path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    suffix = file_path.suffix.lower()

    if suffix == '.docx':
        items = extract_from_docx(file_path)
    elif suffix == '.xlsx':
        items = extract_from_xlsx(file_path)
    elif suffix == '.pdf':
        items = extract_from_pdf(file_path)
    else:
        raise ValueError(f"Unsupported file type: {suffix}")

    saved_paths = []
    for item in items:
        path = save_item(item, output_dir, item_type)
        saved_paths.append(path)
        print(f"Saved: {path}")

    return saved_paths


def import_to_domain(domain: str, file_path: str, item_type: str,
                     config: Optional[dict] = None) -> list[Path]:
    """Import data from a file to a specific domain."""
    if config is None:
        config = get_domain_config(domain)

    base_dir = Path(config['database_path'])
    type_map = {'product': 'products', 'customer': 'customers', 'operator': 'operators'}
    output_dir = base_dir / type_map[item_type]

    return import_file(file_path, str(output_dir), item_type)
